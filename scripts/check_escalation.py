#!/usr/bin/env python3
"""Step 8's Done-when conditions, as assertions.

    scripts/check_escalation.py                  every fixture seed, reseeding each
    scripts/check_escalation.py --seed 42 -v     one seed, against the database as it stands
    scripts/check_escalation.py --keep /tmp/out  leave the packs and workbooks behind

check_pack.py asserts that the right engagements were looked at. This asserts what was said
about them: that mess case 4 comes out NEEDS REVIEW rather than RED, that case 8 is flagged
for continuity with no recommendation attached, that an engagement past its ceiling is
reported over budget with the denominator untouched, and that case 3's silence is never read
as completion.

It reuses check_pack.gather() so the pack is the one the skill's order of operations produces,
then runs classify.py exactly as the skill instructs and reads the exceptions back off the
Exceptions tab of a real workbook rather than out of the dict that produced it.

Three assertions here are not Done-when conditions and are the ones most likely to catch a
regression:

    the classifier's median fee equals percentile_cont(0.5) in mess_cases.sql, per seed
    no unexamined engagement would have been flagged, checked by pulling the financials
        step 4 deliberately did not pull
    the same pack classified twice is byte-identical

The second is the interesting one. scope-escalation can only judge what assemble-delivery-pack
chose to examine, so the four triage triggers have to be sufficient for this policy as well as
for their own. Nothing else in the build checks that, and it is a silent failure if it breaks:
the pack would simply not mention an engagement.
"""

from __future__ import annotations

import argparse
import asyncio
import json
import logging
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any

REPO_ROOT = Path(__file__).resolve().parents[1]
SKILL_DIR = REPO_ROOT / "plugin" / "skills" / "scope-escalation"
sys.path.insert(0, str(REPO_ROOT))
sys.path.insert(0, str(REPO_ROOT / "scripts"))
sys.path.insert(0, str(SKILL_DIR / "scripts"))

# The bundled script, imported as well as shelled out to. The subprocess runs are what the
# skill documents and are what every assertion reads; the import is only for the two audits
# that need to ask "what would this have decided" about data the pack does not contain.
import classify as classifier  # noqa: E402

from check_pack import BUILDER, gather, sheet_rows  # noqa: E402
from check_tools import PERIOD, Checks, anchors_for, report, reseed  # noqa: E402
from mcp import Client  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from server import db, toollog  # noqa: E402
from server.app import mcp  # noqa: E402

CLASSIFIER = SKILL_DIR / "scripts" / "classify.py"
SKILL = SKILL_DIR / "SKILL.md"

FIXTURE_SEEDS = (42, 43, *range(9001, 9016))

EXCEPTION_COLUMNS = (
    "engagement_id", "flag", "triggers", "situation", "cause", "recommended_action",
    "decision_owner",
)

RED = "RED"
NEEDS_REVIEW = "NEEDS REVIEW"

# Language that moves a ceiling. The whole rule is that the denominator never changes, so the
# words for changing it may not appear next to one.
REBASELINE = (
    "re-baseline", "rebaseline", "re-baselined", "rebaselined", "revised ceiling",
    "adjusted ceiling", "reset the ceiling", "updated ceiling", "new ceiling",
    "increase the ceiling", "raise the ceiling", "revised budget", "adjusted budget",
)
REBASELINE_RE = re.compile("|".join(re.escape(w) for w in REBASELINE), re.IGNORECASE)

# Language that reads absent data as a finished engagement. Word boundaries, so
# "data_completeness" does not trip "complete".
COMPLETION = (
    "complete", "completed", "finished", "wrapped up", "concluded", "delivered in full",
    "closed out", "no longer active", "work has ended",
)
COMPLETION_RE = re.compile(r"\b(" + "|".join(re.escape(w) for w in COMPLETION) + r")\b", re.IGNORECASE)


def classify(pack_path: Path, out: Path) -> subprocess.CompletedProcess[str]:
    """Run the bundled classifier the way the skill documents it."""
    return subprocess.run(
        [sys.executable, str(CLASSIFIER), str(pack_path), "--out", str(out)],
        capture_output=True, text=True,
    )


def by_id(rows: list[dict[str, Any]]) -> dict[int, dict[str, Any]]:
    return {int(r["engagement_id"]): r for r in rows}


def flowed(text: str) -> str:
    """Markdown prose with its line wrapping collapsed.

    The phrase assertions below match multi-word sentences against a hard-wrapped document,
    where a phrase is as likely as not to straddle a newline. Matching the raw text makes
    reflowing a paragraph a test failure, which teaches the next person to stop reflowing
    paragraphs rather than to keep the rule.
    """
    return re.sub(r"\s+", " ", text)


# ------------------------------------------------------------------ the Done-when four


def check_case_4(exceptions: dict[int, dict[str, Any]], anchors: dict[str, Any],
                 engagements: dict[int, dict[str, Any]], financials: dict[str, Any],
                 checks: Checks) -> None:
    """The contradictory engagement: fixed fee, under water, burn inside the ceiling.

    It fires the spec's fixed-fee RED rule on margin alone. The precedence rule makes it
    NEEDS REVIEW, and this is the assertion that the precedence rule is actually wired up
    rather than merely written down.
    """
    eid = anchors["case_4_fixed_fee_engagement"]
    row = exceptions.get(eid)
    fin = financials.get(str(eid), {})
    burn_pct = engagements.get(eid, {}).get("burn_pct")

    checks.add(
        "mess case 4: the contradictory engagement is flagged at all",
        row is not None,
        f"engagement {eid}, margin {fin.get('margin_pct')}%, burn {burn_pct}%",
    )
    if row is None:
        return

    checks.add(
        "mess case 4: flagged NEEDS REVIEW rather than RED, though a RED trigger fired",
        row["flag"] == NEEDS_REVIEW and "fixed fee below the margin floor" in row["triggers"],
        f"{row['flag']}; triggers: {row['triggers']}",
    )
    checks.add(
        "mess case 4: the situation states both figures rather than one",
        str(fin.get("margin_pct")) in row["situation"] and str(burn_pct) in row["situation"],
        row["situation"][:96],
    )
    checks.add(
        "mess case 4: the contradiction is stated and left unresolved",
        "disagree" in row["situation"] and "does not say which" in row["situation"],
        "stated as a disagreement with no resolution",
    )
    checks.add(
        "mess case 4: no cause was invented for it",
        not str(row.get("cause") or "").strip()
        or row["cause"] == "not determinable from available data",
        f"cause: {row.get('cause') or 'empty'}",
    )


def check_case_8(exceptions: dict[int, dict[str, Any]], anchors: dict[str, Any],
                 engagements: dict[int, dict[str, Any]], checks: Checks) -> None:
    """Key-person concentration: flagged for continuity, with nothing recommended."""
    eid = anchors["case_8_concentration_engagement"]
    row = exceptions.get(eid)
    engagement = engagements.get(eid, {})

    checks.add(
        "mess case 8: the concentrated engagement is flagged",
        row is not None,
        f"engagement {eid} at {engagement.get('person_concentration_pct')}% "
        f"concentration, band {engagement.get('health_band')}, burn {engagement.get('burn_pct')}%",
    )
    if row is None:
        return

    # Concentration alone is never RED. It can share a row with a trigger that is, though —
    # on seed 9009 the same engagement also carries a high-confidence projected overrun, and
    # RED is right there. What the rule forbids is escalating on the concentration itself.
    triggers = {t.strip() for t in str(row["triggers"]).split(";")}
    concentration_only = triggers == {"concentration above median fee"}

    checks.add(
        "mess case 8: flagged on concentration, and NEEDS REVIEW unless something else is RED",
        "concentration above median fee" in triggers
        and (row["flag"] == NEEDS_REVIEW or not concentration_only),
        f"{row['flag']}; triggers: {row['triggers']}",
    )
    if concentration_only:
        checks.add(
            "mess case 8: no recommendation is attached to the continuity flag",
            row["recommended_action"] == "None. This trigger is a reason to look, not a decision to take.",
            f"recommended_action: {str(row['recommended_action'])[:64]}",
        )


def check_ceiling_never_moves(exceptions: dict[int, dict[str, Any]],
                              engagements: dict[int, dict[str, Any]], checks: Checks) -> None:
    """Over-ceiling engagements report over 100% burn against the ceiling they were given."""
    over = {
        eid: e for eid, e in engagements.items()
        if e.get("ceiling_hours") and e.get("hours_to_date")
        and float(e["hours_to_date"]) > float(e["ceiling_hours"])
    }
    checks.add(
        "over-ceiling engagements report burn above 100%, not a rescaled figure",
        all(float(e["burn_pct"]) > 100 for e in over.values()),
        f"{len(over)} engagement(s) over ceiling"
        + (f", burn {sorted(round(float(e['burn_pct']), 1) for e in over.values())}" if over else ""),
    )

    flagged_over = {eid: exceptions[eid] for eid in over if eid in exceptions}
    checks.add(
        "the ceiling quoted in an over-budget situation is the tool's ceiling",
        all(
            f"{float(engagements[eid]['ceiling_hours']):,.1f}" in row["situation"]
            or "ceiling" not in row["situation"]
            for eid, row in flagged_over.items()
        ),
        f"{len(flagged_over)} over-ceiling engagement(s) flagged",
    )


def check_case_3(exceptions: dict[int, dict[str, Any]], anchors: dict[str, Any],
                 burn: dict[str, Any], checks: Checks) -> None:
    """The silent engagement is flagged, and never read as finished."""
    eid = anchors["case_3_silent_engagement"]
    row = exceptions.get(eid)
    days = (burn.get(str(eid)) or {}).get("days_since_last_entry")

    checks.add(
        "mess case 3: the silent engagement is flagged rather than assumed complete",
        row is not None,
        f"engagement {eid}, {days} days since the last entry",
    )
    if row is None:
        return

    checks.add(
        "mess case 3: flagged RED on the fourteen-day silence",
        row["flag"] == RED and "no time logged for 14 days" in row["triggers"],
        f"{row['flag']}; triggers: {row['triggers']}",
    )
    checks.add(
        "mess case 3: the situation says the data is absent, not that the work is done",
        "Absent data is not evidence of completion." in row["situation"],
        row["situation"][:96],
    )


# ------------------------------------------------------------------ wording, over the file


def check_wording(exceptions: list[dict[str, Any]], checks: Checks) -> None:
    """Two scans over every string that reached the tab.

    Run over the workbook rather than the classifier's output, because the agent writes two
    of the seven columns and the classifier cannot police those. Here the agent is this
    harness, which writes nothing, so what this proves is that the scan works and that the
    script's own prose is clean. The Cowork run is where it has something to catch.
    """
    text = " ".join(
        str(value) for row in exceptions for value in row.values() if value is not None
    )

    rebaselined = REBASELINE_RE.findall(text)
    checks.add(
        "nothing on the Exceptions tab describes a ceiling as moved",
        not rebaselined,
        "no re-baselining language" if not rebaselined else f"found {sorted(set(rebaselined))}",
    )

    completed = COMPLETION_RE.findall(text)
    checks.add(
        "nothing on the Exceptions tab reads absent data as a finished engagement",
        not completed,
        "no completion language" if not completed else f"found {sorted(set(completed))}",
    )


def check_structure(exceptions: list[dict[str, Any]], engagements: dict[int, dict[str, Any]],
                    checks: Checks) -> None:
    """The shape of every row, whatever the seed threw up."""
    checks.add(
        "every exception carries all seven columns",
        all(set(EXCEPTION_COLUMNS) <= set(row) for row in exceptions),
        f"{len(exceptions)} row(s), {len(EXCEPTION_COLUMNS)} columns each",
    )
    checks.add(
        "every flag is RED or NEEDS REVIEW and nothing else",
        all(row["flag"] in (RED, NEEDS_REVIEW) for row in exceptions),
        f"{sum(1 for r in exceptions if r['flag'] == RED)} RED, "
        f"{sum(1 for r in exceptions if r['flag'] == NEEDS_REVIEW)} NEEDS REVIEW",
    )
    checks.add(
        "the decision owner is the engagement lead on every row, never the agent",
        all(row["decision_owner"] == "engagement lead" for row in exceptions),
        f"{len(exceptions)} row(s)",
    )
    checks.add(
        "every flagged engagement names a trigger",
        all(str(row["triggers"]).strip() for row in exceptions),
        "all rows carry their triggers",
    )
    checks.add(
        "every flagged engagement is one that exists in the portfolio",
        all(int(row["engagement_id"]) in engagements for row in exceptions),
        f"{len(exceptions)} of {len(engagements)} engagements flagged",
    )


# ------------------------------------------------------- the three that are not Done-when


def check_median_matches_sql(pack: dict[str, Any], checks: Checks) -> None:
    """The classifier's median fee against percentile_cont(0.5), which mess_cases.sql uses.

    Two definitions of a median disagree on every even-sized portfolio, and the active book
    is eighteen. This is the one number classify.py computes, so it is the one that can drift.
    Goes through db.backend() rather than psycopg so it asks the same question on Aurora.
    """
    ours = classifier.median_fee(pack["engagements"])
    theirs = float(db.backend().query(
        "select percentile_cont(0.5) within group (order by ceiling_amount) as fee "
        "from engagements where status = 'active'"
    )[0]["fee"])

    checks.add(
        "the classifier's median fee is the one mess_cases.sql uses",
        ours is not None and abs(ours - theirs) < 0.01,
        f"classifier {ours:,.2f}, percentile_cont {theirs:,.2f}",
    )


def check_triage_is_sufficient(pack: dict[str, Any], client_rows: dict[str, Any],
                               checks: Checks) -> None:
    """No unexamined engagement would have been flagged had it been examined.

    scope-escalation only sees what assemble-delivery-pack pulled detail for. That is a
    designed boundary and it is only safe if the four triage triggers are wider than this
    skill's, so this pulls the financials step 4 deliberately skipped and checks.
    """
    fee_median = classifier.median_fee(pack["engagements"])
    unexamined = [e for e in pack["engagements"] if str(e["engagement_id"]) not in pack["burn"]]

    would_flag = []
    for engagement in unexamined:
        eid = str(engagement["engagement_id"])
        row = classifier.classify_one(
            engagement, client_rows["burn"].get(eid), client_rows["financials"].get(eid), fee_median
        )
        if row is not None:
            would_flag.append((engagement["engagement_id"], row["flag"], row["triggers"]))

    checks.add(
        "no engagement the triage step skipped would have been flagged",
        not would_flag,
        f"{len(unexamined)} unexamined, none flagged" if not would_flag else f"missed {would_flag}",
    )


def check_determinism(pack_path: Path, workdir: Path, checks: Checks) -> None:
    """The same pack classified twice, byte for byte."""
    first, second = workdir / "determinism-a.json", workdir / "determinism-b.json"
    classify(pack_path, first)
    classify(pack_path, second)
    checks.add(
        "classifying the same pack twice produces the same bytes",
        first.exists() and second.exists() and first.read_bytes() == second.read_bytes(),
        f"{first.stat().st_size if first.exists() else 0} bytes each",
    )


def check_clean_pack(pack: dict[str, Any], workdir: Path, checks: Checks) -> None:
    """A portfolio with nothing wrong flags nothing, and the tab still says so.

    Synthesised rather than found: no fixture seed produces a clean book, and a rule that has
    never been exercised on the empty case is a rule that fails the first quiet month.
    """
    clean = json.loads(json.dumps(pack, default=str))
    for engagement in clean["engagements"]:
        engagement["burn_pct"] = 40.0
        engagement["person_concentration_pct"] = 30.0
    for burn in clean["burn"].values():
        burn["projected_overrun_pct"] = 0.0
        burn["projection_confidence"] = "high"
        burn["days_since_last_entry"] = 1
    for financials in clean["financials"].values():
        financials["margin_pct"] = 42.0

    path = workdir / "clean.json"
    path.write_text(json.dumps(clean, default=str))
    out = workdir / "clean-classified.json"
    result = classify(path, out)

    if result.returncode != 0:
        checks.add("a clean portfolio classifies without error", False, result.stderr[:96])
        return

    exceptions = json.loads(out.read_text())["exceptions"]
    checks.add(
        "a portfolio with nothing wrong produces no exceptions",
        exceptions == [],
        f"{len(exceptions)} exception(s)",
    )

    book = workdir / "clean.xlsx"
    subprocess.run(
        [sys.executable, str(BUILDER), str(out), "--out", str(book)],
        capture_output=True, text=True, check=False,
    )
    sentinel = load_workbook(book)["Exceptions"].cell(row=2, column=1).value if book.exists() else None
    checks.add(
        "an empty Exceptions tab still says nothing was flagged",
        sentinel == "No engagement was flagged by scope-escalation this period.",
        str(sentinel)[:72],
    )


def check_refusals(pack: dict[str, Any], workdir: Path, checks: Checks) -> None:
    """The classifier stops rather than returning an empty list that means 'not looked at'."""
    for name, mutate, expected in (
        ("a pack with no engagements", lambda p: p.update(engagements=[]), "engagements is empty"),
        ("a pack with no burn detail", lambda p: p.update(burn={}), "burn is empty"),
        ("a pack with no financials key", lambda p: p.pop("financials"), "has no 'financials'"),
    ):
        bad = json.loads(json.dumps(pack, default=str))
        mutate(bad)
        path = workdir / "bad-classify.json"
        path.write_text(json.dumps(bad, default=str))
        result = classify(path, workdir / "bad-classify-out.json")
        checks.add(
            f"classify.py refuses {name}",
            result.returncode != 0 and expected in result.stderr,
            result.stderr.strip().splitlines()[-1][:96] if result.stderr.strip() else "no message",
        )


def check_skill_document(checks: Checks) -> None:
    """The skill is the deliverable. Its exact wording is what the run reads."""
    raw = SKILL.read_text() if SKILL.exists() else ""
    checks.add("skill: SKILL.md exists", bool(raw), str(SKILL.relative_to(REPO_ROOT)))
    if not raw:
        return

    text = flowed(raw)
    front = raw.split("---")[1] if raw.startswith("---") else ""
    checks.add(
        "skill: frontmatter names the skill and describes when to use it",
        "name: scope-escalation" in front and "description:" in front and len(front) > 300,
        f"{len(front)} chars of frontmatter",
    )

    for phrase, what in (
        ("projected_overrun_pct > 10", "the overrun trigger"),
        ("margin_pct < 15", "the fixed-fee margin floor"),
        ("burn_pct > 90", "the burn ceiling trigger"),
        ("days_since_last_entry >= 14", "the fourteen-day silence trigger"),
        ("person_concentration_pct > 70", "the concentration trigger"),
        ("above the median fee", "the median-fee qualifier on concentration"),
        ("A contradiction outranks a RED trigger", "the precedence rule, as a heading"),
        ("Never re-baseline a ceiling", "the ceiling rule"),
        ("not determinable from available data", "the cause refusal"),
        ("Never change a flag the script assigned", "the do-not-edit-the-flags rule"),
        ("stop and say so", "the stop rule"),
        ("Do not supply its judgment yourself", "the missing-skill rule"),
    ):
        checks.add(f"skill: states {what}", phrase in text, "present" if phrase in text else "absent")

    owner = "the engagement lead. Never the agent"
    checks.add(
        "skill: puts the decision on the engagement lead, in those words",
        owner in text,
        "stated" if owner in text else "absent",
    )


def check_sibling_skills(checks: Checks) -> None:
    """Step 6 and 7's carried item: a missing skill has to stop a run the way a script does."""
    rule = "Do not supply its judgment yourself"
    for name in ("assemble-delivery-pack", "house-format"):
        path = REPO_ROOT / "plugin" / "skills" / name / "SKILL.md"
        text = flowed(path.read_text()) if path.exists() else ""
        checks.add(
            f"{name}: stops on a missing skill, not only a missing script",
            rule in text,
            "rule present" if rule in text else "absent",
        )


# --------------------------------------------------------------------------- one seed


async def run_seed(seed: int, anchors: dict[str, Any], checks: Checks, workdir: Path) -> None:
    toollog.logger.handlers.clear()
    toollog.logger.addHandler(logging.NullHandler())
    toollog.logger.propagate = False

    period_end = anchors.get("period_end") or "2026-08-31"
    check_skill_document(checks)
    check_sibling_skills(checks)

    async with Client(mcp) as client:
        pack = await gather(client, f"escalation-check-seed-{seed}", period_end)

        # Detail for the engagements triage skipped, used only by the sufficiency audit. The
        # pack itself stays exactly as the skill's order of operations produced it.
        every: dict[str, Any] = {"burn": {}, "financials": {}}
        for engagement in pack["engagements"]:
            eid = str(engagement["engagement_id"])
            if eid in pack["burn"]:
                every["burn"][eid] = pack["burn"][eid]
                every["financials"][eid] = pack["financials"][eid]
                continue
            for tool, key, kwargs in (
                ("get_engagement_burn", "burn", {"as_of_date": period_end}),
                ("get_financials", "financials", {"period": PERIOD}),
            ):
                result = await client.call_tool(
                    tool, {"run_id": f"escalation-audit-seed-{seed}",
                           "engagement_id": int(eid), **kwargs}
                )
                every[key][eid] = result.structured_content

    pack_path = workdir / f"pack-{seed}.json"
    pack_path.write_text(json.dumps(pack, default=str, indent=2))

    classified_path = workdir / f"pack-{seed}-classified.json"
    result = classify(pack_path, classified_path)
    checks.add(
        "classify.py ran and wrote the exceptions",
        result.returncode == 0 and classified_path.exists(),
        result.stdout.strip()[:96] if result.returncode == 0 else result.stderr.strip()[:96],
    )
    if not classified_path.exists():
        return

    classified = json.loads(classified_path.read_text())

    # Everything below reads the workbook, not the dict that produced it.
    book = workdir / f"engagement-book-{PERIOD}-seed-{seed}.xlsx"
    built = subprocess.run(
        [sys.executable, str(BUILDER), str(classified_path), "--out", str(book)],
        capture_output=True, text=True,
    )
    checks.add(
        "the flagged engagements reach a real workbook",
        built.returncode == 0 and book.exists(),
        built.stdout.strip()[:96] if built.returncode == 0 else built.stderr.strip()[:96],
    )
    if not book.exists():
        return

    # openpyxl reads an empty cell back as None, and the classifier writes "" into two of the
    # seven columns on purpose. Normalising here keeps every assertion below reading strings.
    header, raw_rows = sheet_rows(load_workbook(book)["Exceptions"])
    rows = [{k: ("" if v is None else v) for k, v in row.items()} for row in raw_rows]
    checks.add(
        "the Exceptions tab carries the triggers column",
        list(header)[:len(EXCEPTION_COLUMNS)] == list(EXCEPTION_COLUMNS),
        ", ".join(header),
    )

    exceptions = [r for r in rows if r.get("flag")]
    checks.add(
        "every exception the classifier wrote reached the tab",
        len(exceptions) == len(classified["exceptions"]),
        f"{len(exceptions)} on the tab, {len(classified['exceptions'])} classified",
    )

    indexed = by_id(exceptions)
    engagements = by_id(pack["engagements"])

    check_case_4(indexed, anchors, engagements, pack["financials"], checks)
    check_case_8(indexed, anchors, engagements, checks)
    check_ceiling_never_moves(indexed, engagements, checks)
    check_case_3(indexed, anchors, pack["burn"], checks)
    check_wording(exceptions, checks)
    check_structure(exceptions, engagements, checks)

    check_median_matches_sql(pack, checks)
    check_triage_is_sufficient(pack, every, checks)
    check_determinism(pack_path, workdir, checks)
    check_clean_pack(pack, workdir, checks)
    check_refusals(pack, workdir, checks)


async def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--seed", type=int, action="append", help="Repeatable. Every fixture seed when omitted.")
    ap.add_argument("--dsn", default=db.DEFAULT_DSN)
    ap.add_argument("--no-reseed", action="store_true", help="Use the database as it stands.")
    ap.add_argument("--verbose", "-v", action="store_true", help="Print passing assertions too.")
    ap.add_argument("--keep", type=Path, help="Write the packs and workbooks here and leave them.")
    ap.add_argument(
        "--backend", choices=("local", "aws"), default="local",
        help="'aws' drives the same sequence over the RDS Data API. Implies --no-reseed.",
    )
    args = ap.parse_args()

    seeds = args.seed or list(FIXTURE_SEEDS)
    if args.no_reseed and len(seeds) > 1:
        ap.error("--no-reseed needs a single --seed, since it cannot change the loaded data")
    if args.backend == "aws" and not args.no_reseed:
        ap.error("--backend aws expects --no-reseed and a single --seed matching what Aurora holds")

    db._backend = db.get_backend(args.backend, args.dsn)

    workdir = args.keep or Path(tempfile.mkdtemp(prefix="bizdata-escalation-"))
    workdir.mkdir(parents=True, exist_ok=True)

    ok = True
    try:
        for seed in seeds:
            if not args.no_reseed:
                reseed(seed, args.dsn)
            checks = Checks(seed)
            try:
                await run_seed(seed, anchors_for(seed), checks, workdir)
            except Exception as exc:  # noqa: BLE001
                while isinstance(exc, BaseExceptionGroup) and exc.exceptions:
                    exc = exc.exceptions[0]
                checks.add("harness ran to completion", False, f"{type(exc).__name__}: {exc}"[:160])
            ok &= report(seed, checks, args.verbose)
    finally:
        if args.keep:
            print(f"\npacks and workbooks left in {workdir}")
        else:
            shutil.rmtree(workdir, ignore_errors=True)

    print()
    print("all seeds pass" if ok else "FAILURES above")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
